from fastapi import (
    APIRouter, Depends, HTTPException, UploadFile, Form, File, status, Request, Query, Body
)
from pydantic import BaseModel, constr
from sqlalchemy import or_, desc,cast, String,asc
from app.mail import send_email
from app.models.user import SupportMessage
from app.schemas.cart import CartItemOut
from app.schemas.cart import CartItemCreate, CartItemOut, CartItemUpdate
from slugify import slugify
import math
from datetime import datetime
from app.schemas.order import OrderOut
from app.schemas.order_system import OrderResponse
from app.schemas.product import BestSellerRequest
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy import func, extract, literal
from sqlalchemy.orm import Session, joinedload
from typing import List, Union, Optional, Dict
import os, shutil, json, time, re
from datetime import datetime, timedelta,date
import pathlib
from app.schemas.product import ProductOut
from app.database import get_db
from app import models, schemas
from app.models import User, Category, Coupon, Order, OrderItem, OrderTracking,WishlistItem
from app.models.product import Product
from app.schemas import category as category_schema
from app.authentication import get_current_admin_user, get_current_user, create_access_token
from app.security import verify_password, get_password_hash  # ✅ FIXED: Import from security instead of authentication
from app.services.password_policy import PasswordPolicyError, validate_password
from app.models.user import PasswordResetToken
from app.crud import user as crud, admin
from app.schemas.user import UserOut
from datetime import datetime
from uuid import uuid4
from fastapi.responses import JSONResponse

import logging

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/admin", tags=["Admin"])

# -------------------------------------------------------------------
#                        GLOBAL / HELPERS / VALIDATORS
# -------------------------------------------------------------------
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent.parent  # root/jokgroupfastapi
STATIC_DIR = BASE_DIR / "static"
PRODUCTS_DIR = STATIC_DIR / "products"
CATEGORIES_DIR = STATIC_DIR / "categories"

PRODUCTS_DIR.mkdir(parents=True, exist_ok=True)
CATEGORIES_DIR.mkdir(parents=True, exist_ok=True)

# Create folders if missing


ALLOWED_IMAGE_CONTENT_TYPES = {
    "image/jpeg",
    "image/png",
    "image/webp",
    "image/avif",
    "image/jfif"
}


ALLOWED_VIDEO_TYPES = ["video/mp4", "video/webm", "video/mkv"]
ALLOWED_PDF_TYPES = ["application/pdf"]

MAX_IMAGE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_IMAGES_PER_PRODUCT = 5
MIN_IMAGES_PER_PRODUCT = 1
MAX_SLUG_LENGTH = 100
MAX_NAME_LENGTH = 255
MAX_DESC_LENGTH = 5000
MAX_CATEGORY_NAME = 150
MAX_SUBCATEGORY_NAME = 150
MAX_VIDEO_SIZE = 50 * 1024 * 1024
MAX_PDF_SIZE = 10 * 1024 * 1024

SAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9_.-]")




def safe_filename(filename: str) -> str:
    if not filename:
        return f"file_{int(time.time())}"
    base = os.path.basename(filename)
    name, ext = os.path.splitext(base)
    safe_name = SAFE_FILENAME_CHARS.sub("_", name)[:120]
    safe_ext = re.sub(r"[^A-Za-z0-9.]", "", ext) or ".jpg"
    return f"{safe_name}{safe_ext}"



def fix_image_url(relative_path: str) -> str:
    """Return a correct URL starting with /static/..."""
    if not relative_path:
        return ""

    # If full URL, return as is
    if relative_path.startswith("http://") or relative_path.startswith("https://"):
        return relative_path

    rel = str(relative_path).lstrip("/")

    # Always prefix with exactly one /static/
    if rel.startswith("static/"):
        rel = rel[len("static/"):]  # remove leading static/

    return f"/static/{rel}"


def _is_image_file(upload: UploadFile) -> bool:
    ct = getattr(upload, "content_type", "") or ""
    if ct.split(";")[0] in ALLOWED_IMAGE_CONTENT_TYPES:
        return True
    ext = (upload.filename or "").split(".")[-1].lower()
    return ext in {"jpg", "jpeg", "png", "webp", "avif","jfif"}


async def validate_and_save_image(upload: UploadFile, dest_filename: str, folder: str = "products") -> str:
    """Save to static/products/ or static/categories/ and return stored path"""
    # Choose save directory
    if folder == "categories":
        save_dir = CATEGORIES_DIR
    else:
        save_dir = PRODUCTS_DIR

    try:
        contents = await upload.read()
    except Exception:
        raise HTTPException(status_code=400, detail="Failed to read uploaded file.")

    size = len(contents)
    if size == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if size > MAX_IMAGE_SIZE_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"Image exceeds maximum size of {MAX_IMAGE_SIZE_BYTES} bytes."
        )
    if not _is_image_file(upload):
        raise HTTPException(status_code=400, detail="Uploaded file is not a permitted image type.")

    save_path = save_dir / dest_filename

    try:
        with open(save_path, "wb") as f:
            f.write(contents)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save image: {str(e)}")

    rel = os.path.relpath(str(save_path), str(STATIC_DIR))  # returns categories/img.jpg or products/img.jpg
    return fix_image_url(rel)


async def save_file_with_limit(
        file: UploadFile,
        allowed_types: list[str],
        max_size: int,
        upload_dir: str  # Add this parameter
) -> str:
    """
    Save uploaded file with type and size validation
    """
    import os
    import time
    from pathlib import Path

    # Validate file type
    if not file.content_type or file.content_type not in allowed_types:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid file type. Allowed: {', '.join(allowed_types)}"
        )

    # Create upload directory if it doesn't exist
    upload_path = Path(upload_dir)
    upload_path.mkdir(parents=True, exist_ok=True)

    # Generate unique filename
    timestamp = int(time.time())
    original_filename = file.filename or "unnamed"
    filename = safe_filename(f"{timestamp}_{original_filename}")
    file_path = upload_path / filename

    # Save file with size validation
    real_file_size = 0

    try:
        with open(file_path, "wb") as buffer:
            while chunk := await file.read(8192):  # Read in 8KB chunks
                real_file_size += len(chunk)

                # Check if file exceeds max size
                if real_file_size > max_size:
                    raise HTTPException(
                        status_code=413,
                        detail=f"File too large. Max size: {max_size / (1024 * 1024):.1f}MB"
                    )

                buffer.write(chunk)

    except HTTPException:
        # Delete partial file if size limit exceeded
        if file_path.exists():
            os.unlink(file_path)
        raise

    finally:
        await file.close()

    # Return relative path for database storage
    return f"/{upload_dir}/{filename}"


def list_to_db_str(values: Optional[Union[List[str], str]]) -> Optional[str]:
    """Convert a list into a comma-separated string for DB storage"""
    if values is None:
        return None
    if isinstance(values, list):
        cleaned = [str(v).strip() for v in values if str(v).strip()]
        return ",".join(cleaned) if cleaned else None
    s = str(values).strip()
    return s if s else None


def safe_parse_list_field(value: Optional[str]) -> List[str]:
    """Accept comma-separated, JSON array string, or list-like string and return Python list."""
    if value is None:
        return []
    if isinstance(value, list):
        return [str(x) for x in value]
    s = str(value).strip()
    if not s:
        return []
    # Try JSON parse
    if s.startswith("[") and s.endswith("]"):
        try:
            parsed = json.loads(s)
            if isinstance(parsed, list):
                return [str(x) for x in parsed]
        except Exception:
            pass
    # comma-split fallback
    return [p.strip() for p in s.split(",") if p.strip()]


def validate_name(name: str, max_len: int = MAX_NAME_LENGTH) -> str:
    if not isinstance(name, str) or not name.strip():
        raise HTTPException(status_code=400, detail="Name is required.")
    v = name.strip()
    if len(v) > max_len:
        raise HTTPException(status_code=400, detail=f"Name too long (max {max_len}).")
    return v


def validate_price(value):
    try:
        v = float(value)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid numeric value.")
    if v < 0:
        raise HTTPException(status_code=400, detail="Value must be non-negative.")
    return v


def validate_rating(r):
    try:
        rv = float(r)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid rating value.")
    if rv < 0 or rv > 5:
        raise HTTPException(status_code=400, detail="Rating must be between 0 and 5.")
    return rv


def validate_reviews_count(rv):
    try:
        iv = int(rv)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid reviews value.")
    if iv < 0:
        raise HTTPException(status_code=400, detail="Reviews must be non-negative.")
    return iv


def ensure_category_exists(db: Session, category_id: Union[int, str]):
    try:
        cid = int(category_id)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid category_id.")
    cat = db.query(models.Category).filter(models.Category.id == cid).first()
    if not cat:
        raise HTTPException(status_code=404, detail="Category not found.")
    return cat


def ensure_subcategory_exists(db: Session, subcategory_id: Union[int, str]):
    try:
        sid = int(subcategory_id)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid subcategory_id.")
    sub = db.query(models.SubCategory).filter(models.SubCategory.id == sid).first()
    if not sub:
        raise HTTPException(status_code=404, detail="Subcategory not found.")
    return sub

# -------------------------------------------------------------------
#                         ADMIN AUTHENTICATION
# -------------------------------------------------------------------

@router.post("/token")
def admin_login(
    form_data: OAuth2PasswordRequestForm = Depends(),
    db: Session = Depends(get_db)
):
    if not form_data.username or not form_data.password:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="username and password are required"
        )

    user = crud.get_user_by_email(db, form_data.username)

    if not user or not verify_password(form_data.password, user.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid credentials"
        )

    if user.role.lower() != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Admin access only"
        )

    access_token = create_access_token(
        data={
            "sub": user.email,
            "role": user.role
        }
    )

    return {
        "access_token": access_token,
        "token_type": "bearer"
    }

@router.get("/me", response_model=UserOut)
def get_admin_me(admin: User = Depends(get_current_admin_user)):
    return admin


@router.get("/users")
def get_all_users(
        db: Session = Depends(get_db),
        admin: User = Depends(get_current_admin_user),

        # Pagination
        page: int = Query(1, ge=1),
        limit: int = Query(20, ge=1, le=100),

        # Search
        search: Optional[str] = Query(None, description="Search by name, email, or phone"),

        # Filters
        role: Optional[str] = Query(None, description="Filter by role"),
        is_active: Optional[bool] = Query(None, description="Filter by active status"),
        created_after: Optional[datetime] = Query(None, description="Filter users created after this date"),
        created_before: Optional[datetime] = Query(None, description="Filter users created before this date"),

        # Sorting
        sort_by: Optional[str] = Query("created_at", description="Sort by field"),
        sort_order: Optional[str] = Query("desc", description="Sort order (asc or desc)"),
):
    """
    List all users (Admin only) with server-side pagination
    """
    query = db.query(User)

    # SEARCH FILTER
    if search:
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                User.first_name.ilike(search_term),
                User.last_name.ilike(search_term),
                User.email.ilike(search_term),
                User.phone.ilike(search_term),
            )
        )

    # ROLE FILTER
    if role:
        query = query.filter(User.role == role)

    # ACTIVE STATUS FILTER
    if is_active is not None:
        query = query.filter(
            User.deleted_at.is_(None) if is_active else User.deleted_at.isnot(None)
        )

    # DATE RANGE FILTERS
    if created_after:
        query = query.filter(User.created_at >= created_after)

    if created_before:
        query = query.filter(User.created_at <= created_before)

    # Get total count before pagination
    total_count = query.count()

    # SORTING
    sort_column = getattr(User, sort_by, User.created_at)
    if sort_order.lower() == "asc":
        query = query.order_by(asc(sort_column))
    else:
        query = query.order_by(desc(sort_column))

    # PAGINATION
    offset = (page - 1) * limit
    users = query.offset(offset).limit(limit).all()

    return {
        "data": users,
        "total": total_count,
        "page": page,
        "limit": limit,
        "total_pages": math.ceil(total_count / limit) if limit > 0 else 0
    }


# Optional: Add stats endpoint
@router.get("/users/stats")
def get_users_stats(
        db: Session = Depends(get_db),
        admin: User = Depends(get_current_admin_user),
):
    """Get user statistics"""
    # The column is `active`, not `is_active` — but count on `deleted_at`:
    # it is what login actually enforces, the activate/deactivate endpoints
    # set it first, and it can't be NULL-ambiguous on rows that predate the
    # `active` boolean. Deriving `inactive` keeps the numbers reconciling.
    total = db.query(User).count()
    active = db.query(User).filter(User.deleted_at.is_(None)).count()
    admins = db.query(User).filter(User.role == 'admin').count()

    return {
        "total": total,
        "active": active,
        "inactive": total - active,
        "admins": admins
    }



# -------------------------------------------------------------------
#                        CATEGORY MANAGEMENT
# -------------------------------------------------------------------

NAME_ALLOWED_RE = re.compile(r"^[A-Za-z0-9\s\-\&\'\.]+$")  # allow letters, numbers, spaces, -, &, ' and dot

def validate_entity_name(value: str, max_len: int):
    """Validate a category/subcategory name: presence, length, allowed chars."""
    if not isinstance(value, str) or not value.strip():
        raise HTTPException(status_code=400, detail="Name is required.")
    v = value.strip()
    if len(v) > max_len:
        raise HTTPException(status_code=400, detail=f"Name too long (max {max_len} chars).")
    if not NAME_ALLOWED_RE.match(v):
        raise HTTPException(
            status_code=400,
            detail="Name contains invalid characters. Allowed: letters, numbers, spaces, hyphen (-), ampersand (&), apostrophe (') and dot (.)"
        )
    return v

# ---------- Replace create_category with this (ensures name validation & uniqueness) ----------
@router.post("/cat/create", response_model=schemas.CategoryOut)
async def create_category(
    name: str = Form(...),
    slug: str = Form(None),
    image: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)
):
    name_clean = validate_entity_name(name, max_len=MAX_CATEGORY_NAME)
    slug_val = slug.strip() if slug else slugify(name_clean)

    # uniqueness checks
    if db.query(models.Category).filter(func.lower(models.Category.name) == name_clean.lower()).first():
        raise HTTPException(status_code=409, detail="Category name already exists.")
    if db.query(models.Category).filter(func.lower(models.Category.slug) == slug_val.lower()).first():
        raise HTTPException(status_code=409, detail="Category slug already exists.")

    # save image
    file_ext = os.path.splitext(image.filename)[-1] or ".jpg"
    filename_safe = safe_filename(f"{slug_val}{file_ext}")
    # validate_and_save_image supports folder="categories"
    image_url = await validate_and_save_image(image, filename_safe, "categories")

    new_category = models.Category(name=name_clean, slug=slug_val, image=image_url)
    db.add(new_category)
    db.commit()
    db.refresh(new_category)

    # return nested structure (no subcategories yet)
    return {
        "id": new_category.id,
        "name": new_category.name,
        "slug": new_category.slug,
        "image": new_category.image,
        "subcategories": []
    }




@router.get("/cat/list", response_model=List[schemas.CategoryOut])
def get_categories(
    db: Session = Depends(get_db),

    # 🔹 Pagination
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),

    # 🔹 Search
    search: Optional[str] = Query(None, description="Search category by name"),
):
    """
    List categories with:
    - Pagination
    - Search by category name
    """

    query = db.query(models.Category)

    # 🔍 SEARCH BY CATEGORY NAME
    if search:
        query = query.filter(models.Category.name.ilike(f"%{search.strip()}%"))

    # ⬇️ ORDER BY NEWEST FIRST
    query = query.order_by(desc(models.Category.id))

    # 📄 PAGINATION
    offset = (page - 1) * limit
    categories = query.offset(offset).limit(limit).all()

    response = []
    for cat in categories:
        subcats = [
            {
                "id": sub.id,
                "name": sub.name or "Unnamed",
                "slug": sub.slug,
                "category_id": sub.category_id,
                "subcategory_id": sub.id,
            }
            for sub in cat.subcategories
        ]

        response.append({
            "id": cat.id,
            "name": cat.name,
            "slug": cat.slug,
            "image": cat.image,
            "is_active": cat.is_active,
            "subcategories": subcats or None,
        })

    return response


class CategoryVisibilityUpdate(BaseModel):
    category_ids: List[int]
    is_active: bool


@router.patch("/cat/visibility")
def set_categories_visibility(
    payload: CategoryVisibilityUpdate,
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Show or hide categories on the storefront in one call.

    Hiding flips is_active only — the category, its subcategories and its
    products stay in the catalogue; the category just vanishes from the
    storefront's menus, listings and its own page until shown again. The
    products remain reachable through All Products and search unless
    hidden individually.
    """
    ids = list({int(i) for i in payload.category_ids})
    if not ids:
        raise HTTPException(status_code=400, detail="No categories selected")
    if len(ids) > 200:
        raise HTTPException(status_code=400, detail="Too many categories in one request (max 200)")

    categories = db.query(models.Category).filter(models.Category.id.in_(ids)).all()
    for category in categories:
        category.is_active = payload.is_active

    db.commit()

    found_ids = [c.id for c in categories]
    missing = sorted(set(ids) - set(found_ids))
    verb = "shown" if payload.is_active else "hidden"
    return {
        "updated": len(found_ids),
        "category_ids": found_ids,
        "missing_ids": missing,
        "is_active": payload.is_active,
        "message": f"{len(found_ids)} {'category' if len(found_ids) == 1 else 'categories'} {verb}",
    }
@router.put("/cat/update/{category_id}", response_model=schemas.CategoryOut)
async def update_category(
    category_id: int,
    name: str = Form(...),
    slug: str = Form(...),
    image: UploadFile = File(None),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)

):
    """Update category"""
    if category_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid category_id")

    category = db.query(models.Category).filter(models.Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")

    name_clean = validate_entity_name(name, max_len=MAX_CATEGORY_NAME)
    slug = validate_name(slug)

    # check duplicate name if changed
    if name_clean.lower() != (category.name or "").lower():
        if db.query(models.Category).filter(func.lower(models.Category.name) == name_clean.lower(), models.Category.id != category_id).first():
            raise HTTPException(status_code=409, detail="Category name already in use by another category.")

    # check slug uniqueness if changed
    if slug.lower() != (category.slug or "").lower():
        if db.query(models.Category).filter(func.lower(models.Category.slug) == slug.lower(), models.Category.id != category_id).first():
            raise HTTPException(status_code=409, detail="Slug already in use by another category.")

    category.name = name_clean
    category.slug = slug

    if image:
        # remove old image file if exists
        if category.image:
            old_image_path = os.path.join(BASE_DIR, "..", "static", category.image.lstrip("/static/"))
            if os.path.exists(old_image_path):
                try:
                    os.remove(old_image_path)
                except Exception:
                    pass

        file_ext = os.path.splitext(image.filename)[-1] or ".jpg"
        file_name = safe_filename(f"{slug}{file_ext}")
        image_path = await validate_and_save_image(image, file_name)
        category.image = image_path

    db.commit()
    db.refresh(category)

    return {
        "id": category.id,
        "name": category.name,
        "slug": category.slug,
        "image": category.image,
        "subcategories": [
            {
                "id": sub.id,
                "name": sub.name,
                "slug": sub.slug,
                "category_id": sub.category_id
            } for sub in category.subcategories
        ]
    }


@router.delete("/cat/delete/{category_id}")
def delete_category(category_id: int, db: Session = Depends(get_db), admin=Depends(get_current_admin_user)):

    """Delete category"""
    if category_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid category_id")

    category = db.query(models.Category).filter(models.Category.id == category_id).first()
    if not category:
        raise HTTPException(status_code=404, detail="Category not found")

    if category.image:
        image_path = os.path.join(BASE_DIR, "..", "static", category.image.lstrip("/static/"))
        if os.path.exists(image_path):
            try:
                os.remove(image_path)
            except Exception:
                pass

    db.delete(category)
    db.commit()
    return {"detail": "Category deleted"}


# -------------------------------------------------------------------
#                        SUBCATEGORY MANAGEMENT
# -------------------------------------------------------------------

@router.post("/cat/subcategory/create", response_model=schemas.SubcategoryOut)
def create_subcategory(
    subcategory: schemas.SubCategoryCreate,
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)
):
    name_clean = validate_entity_name(subcategory.name, max_len=MAX_SUBCATEGORY_NAME)
    slug_val = subcategory.slug.strip() if subcategory.slug else slugify(name_clean)

    parent = ensure_category_exists(db, subcategory.category_id)

    # uniqueness within category
    if db.query(models.SubCategory).filter(
        func.lower(models.SubCategory.name) == name_clean.lower(),
        models.SubCategory.category_id == parent.id
    ).first():
        raise HTTPException(status_code=409, detail="Subcategory name already exists in this category.")

    if db.query(models.SubCategory).filter(
        func.lower(models.SubCategory.slug) == slug_val.lower(),
        models.SubCategory.category_id == parent.id
    ).first():
        raise HTTPException(status_code=409, detail="Subcategory slug already exists in this category.")

    new_sub = models.SubCategory(name=name_clean, slug=slug_val, category_id=parent.id)
    db.add(new_sub)
    db.commit()
    db.refresh(new_sub)

    return {
        "id": new_sub.id,
        "name": new_sub.name,
        "slug": new_sub.slug,
        "category_id": new_sub.category_id,
        "subcategory_id": new_sub.id
    }

@router.get("/cat/subcategory/list", response_model=List[schemas.SubcategoryOut])
def get_all_subcategories(
    db: Session = Depends(get_db),

    # 🔹 Pagination
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),

    # 🔹 Search
    search: Optional[str] = Query(None, description="Search subcategory by name"),
):
    """
    List subcategories with:
    - Pagination
    - Search by subcategory name
    """

    query = db.query(models.SubCategory)

    # 🔍 SEARCH BY NAME
    if search:
        query = query.filter(models.SubCategory.name.ilike(f"%{search.strip()}%"))

    # ⬇️ ORDER BY NEWEST FIRST
    query = query.order_by(desc(models.SubCategory.id))

    # 📄 PAGINATION
    offset = (page - 1) * limit
    subs = query.offset(offset).limit(limit).all()

    return [
        {
            "id": s.id,
            "name": s.name,
            "slug": s.slug,
            "category_id": s.category_id,
            "subcategory_id": s.id,
        }
        for s in subs
    ]

@router.put("/cat/sub/update/{sub_id}", response_model=schemas.SubcategoryOut)
def update_subcategory(sub_id: int, sub_data: schemas.SubCategoryUpdate, db: Session = Depends(get_db), admin=Depends(get_current_admin_user)):
    """Update subcategory"""
    if sub_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid sub_id")
    sub = db.query(models.SubCategory).filter(models.SubCategory.id == sub_id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="SubCategory not found")

    name_clean = validate_entity_name(sub_data.name, max_len=MAX_SUBCATEGORY_NAME)
    slug = validate_name(sub_data.slug)

    # check name uniqueness within same category if changed
    if name_clean.lower() != (sub.name or "").lower():
        if db.query(models.SubCategory).filter(func.lower(models.SubCategory.name) == name_clean.lower(), models.SubCategory.category_id == sub.category_id, models.SubCategory.id != sub_id).first():
            raise HTTPException(status_code=409, detail="Subcategory name already exists in this category.")

    # check slug uniqueness within same category if changed
    if slug.lower() != (sub.slug or "").lower():
        if db.query(models.SubCategory).filter(func.lower(models.SubCategory.slug) == slug.lower(), models.SubCategory.category_id == sub.category_id, models.SubCategory.id != sub_id).first():
            raise HTTPException(status_code=409, detail="Subcategory slug already exists in this category.")

    sub.name = name_clean
    sub.slug = slug
    db.commit()
    db.refresh(sub)
    return {
        "id": sub.id,
        "name": sub.name,
        "slug": sub.slug,
        "category_id": sub.category_id,
        "subcategory_id": sub.id
    }
@router.delete("/cat/sub/delete/{sub_id}")
def delete_subcategory(sub_id: int, db: Session = Depends(get_db) ,admin=Depends(get_current_admin_user)):
    """Delete subcategory"""
    if sub_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid sub_id")
    sub = db.query(models.SubCategory).filter(models.SubCategory.id == sub_id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="SubCategory not found")

    db.delete(sub)
    db.commit()
    return {"detail": "SubCategory deleted"}


# -------------------------------------------------------------------
#                        PRODUCT MANAGEMENT
# -------------------------------------------------------------------

@router.post("/product/create-new-arrival")
async def create_new_arrival_product(
    name: str = Form(...),
    description: str = Form(...),
    price: float = Form(...),
    discount_price: float = Form(0),
    category_id: str = Form(...),
    subcategory_id: str = Form(...),
    colors: str = Form(None),
    sizes: str = Form(None),
    in_stock: bool = Form(True),
    rating: float = Form(0.0),
    reviews: int = Form(0),
    images: List[UploadFile] = File(...),
    db: Session = Depends(get_db),
    current_admin: User = Depends(get_current_admin_user)
):
    """Create new arrival product with validation"""
    name = validate_name(name)
    if len(name) > MAX_NAME_LENGTH:
        raise HTTPException(status_code=400, detail=f"Product name too long (max {MAX_NAME_LENGTH}).")
    description = description.strip() if isinstance(description, str) else ""
    if not description:
        raise HTTPException(status_code=400, detail="Description is required.")
    if len(description) > MAX_DESC_LENGTH:
        raise HTTPException(status_code=400, detail=f"Description too long (max {MAX_DESC_LENGTH} chars).")

    price = validate_price(price)
    discount_price = validate_price(discount_price)
    if discount_price > price:
        raise HTTPException(status_code=400, detail="Discount price cannot exceed price.")

    # validate category and subcategory existence and relation
    category = ensure_category_exists(db, category_id)
    subcategory = ensure_subcategory_exists(db, subcategory_id)
    if getattr(subcategory, "category_id", None) != category.id:
        raise HTTPException(status_code=400, detail="Subcategory does not belong to provided category.")

    # images validation
    if not images or len(images) < MIN_IMAGES_PER_PRODUCT:
        raise HTTPException(status_code=400, detail=f"At least {MIN_IMAGES_PER_PRODUCT} image(s) required.")
    if len(images) > MAX_IMAGES_PER_PRODUCT:
        raise HTTPException(status_code=400, detail=f"Maximum {MAX_IMAGES_PER_PRODUCT} images allowed.")
    image_path = []
    for image in images:
        if not getattr(image, "filename", None):
            raise HTTPException(status_code=400, detail="Each image must have a filename.")
        filename = safe_filename(f"{int(time.time())}_{image.filename}")
        saved_path = await validate_and_save_image(image, filename)
        image_path.append(saved_path)

    parsed_colors = safe_parse_list_field(colors)
    parsed_sizes = safe_parse_list_field(sizes)

    # optional: ensure product name uniqueness (case-insensitive)
    if db.query(Product).filter(func.lower(Product.name) == name.lower()).first():
        raise HTTPException(status_code=400, detail="Product with this name already exists.")

    new_product = Product(
        name=name,
        description=description,
        price=price,
        discount_price=discount_price,
        category_id=category.id,
        subcategory_id=subcategory.id,
        sizes=parsed_sizes,  # ✅ No json.dumps
        in_stock=bool(in_stock),
        rating=validate_rating(rating),
        reviews=validate_reviews_count(reviews),
        featured=False,
        best_seller=False,
        new_arrival=True,
    )
    db.add(new_product)
    db.commit()
    db.refresh(new_product)

    return {"message": "New Arrival product created successfully!", "product_id": new_product.id}


@router.get("/product/new-arrivals")
def get_new_arrivals(
    db: Session = Depends(get_db),

    # 🔹 Pagination
    page: int = Query(1, ge=1),
    limit: int = Query(200, ge=1, le=100),
):
    """List new arrival products with pagination"""

    def parse_images(value):
        if not value:
            return []
        try:
            if isinstance(value, str) and value.strip().startswith("["):
                return json.loads(value)
        except Exception:
            pass
        if isinstance(value, str) and "," in value:
            return [img.strip() for img in value.split(",") if img.strip()]
        return [value] if isinstance(value, str) else []

    # ⬇️ newest first
    query = (
        db.query(Product)
        .filter(Product.new_arrival == True)
        .order_by(Product.id.desc())
    )

    # 📄 pagination
    offset = (page - 1) * limit
    products = query.offset(offset).limit(limit).all()

    # 📄 PAGINATION

    def to_list(value):
        if isinstance(value, str):
            return [v.strip() for v in value.split(",") if v.strip()]
        elif isinstance(value, list):
            return [v.strip() for v in value if isinstance(v, str) and v.strip()]
        return []

    result = []
    for p in products:
        result.append({
            "id": p.id,
            "name": p.name,
            "description": p.description,
            "details": p.details,
            "specifications": p.specifications,
            "highlights": to_list(p.highlights),
            "price": p.price,
            "discount_price": p.discount_price,
            "category_id": p.category_id,
            "subcategory_id": p.subcategory_id,
            "category": p.category.name if p.category else None,
            "subcategory": p.subcategory.name if p.subcategory else None,
            "sizes": to_list(p.sizes),
            "colors": to_list(p.colors),
            "images": to_list(p.images),
            "video_url": p.video_url,
            "pdf_url": p.pdf_url,
            "in_stock": p.in_stock,
            "is_active": p.is_active,
            "hsn_code": p.hsn_code,
            "sgst": p.sgst,
            "cgst": p.cgst,
            "rating": p.rating,
            "reviews": p.reviews,
            "featured": p.featured,
            "best_seller": p.best_seller,
            "new_arrival": p.new_arrival,
            "created_at": p.created_at,
        })

    return result


def _validate_gst_percent(value: Optional[float], label: str) -> Optional[float]:
    """SGST/CGST are percentages — half of the product's GST rate each."""
    if value is None:
        return None
    try:
        v = float(value)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{label} must be a number")
    if not (0 <= v <= 100):
        raise HTTPException(status_code=400, detail=f"{label} must be between 0 and 100")
    return round(v, 2)


def _ensure_list_from_form(val: Optional[str]):
    if not val:
        return []
    val = val.strip()
    # JSON array?
    if (val.startswith("[") and val.endswith("]")) or (val.startswith("{") and val.endswith("}")):
        try:
            parsed = json.loads(val)
            if isinstance(parsed, list):
                return parsed
            if isinstance(parsed, str):
                return [parsed]
        except Exception:
            pass
    if "," in val:
        return [s.strip() for s in val.split(",") if s.strip()]
    return [val]

async def _save_file_to_disk(upload: UploadFile, dest_path: str) -> None:
    """Write uploaded file to dest_path (async wrapper around sync I/O)."""
    def _write():
        with open(dest_path, "wb") as out_f:
            shutil.copyfileobj(upload.file, out_f)
    # Note: to_thread might not be available, using direct sync call for now
    _write()
    # ensure file pointer closed
    try:
        upload.file.close()
    except Exception:
        pass

def _build_image_url(product_id: int, filename: str) -> str:
    return f"/static/products/{product_id}/{filename}"

def _product_files_dir(product_id: int) -> str:
    return os.path.join(PRODUCTS_DIR, str(product_id))

# ---------------- CREATE ----------------
@router.post("/product/create", response_model=ProductOut, status_code=201)
async def create_product(
    name: str = Form(...),
    description: Optional[str] = Form(None),
    price: float = Form(...),
    discount_price: Optional[float] = Form(None),
    category_id: int = Form(...),
    subcategory_id: Optional[int] = Form(None),
    colors: Optional[str] = Form(None),
    sizes: Optional[str] = Form(None),
    images: Optional[List[UploadFile]] = File(None),
    video: Optional[UploadFile] = File(None),
    pdf: Optional[UploadFile] = File(None),
    in_stock: bool = Form(True),
    rating: Optional[float] = Form(None),
    reviews: Optional[int] = Form(0),
    featured: bool = Form(False),
    best_seller: bool = Form(False),
    new_arrival: bool = Form(False),
    highlights: Optional[str] = Form(None),
    specifications: Optional[str] = Form(None),
    details: Optional[str] = Form(None),
    sgst: Optional[float] = Form(None),
    cgst: Optional[float] = Form(None),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):

    # ---------------- VALIDATIONS ----------------
    name = validate_name(name)
    price = validate_price(price)
    sgst = _validate_gst_percent(sgst, "SGST")
    cgst = _validate_gst_percent(cgst, "CGST")

    if discount_price is not None:
        discount_price = validate_price(discount_price)
        if discount_price > price:
            raise HTTPException(400, "Discount price cannot exceed price")

    if rating is not None:
        rating = validate_rating(rating)

    reviews = validate_reviews_count(reviews or 0)

    category = ensure_category_exists(db, category_id)

    if subcategory_id:
        sub = ensure_subcategory_exists(db, subcategory_id)
        if sub.category_id != category.id:
            raise HTTPException(400, "Subcategory does not belong to category")
    else:
        sub = None

    # ---------------- IMAGE UPLOAD ----------------
    image_urls: list[str] = []

    if images:
        if len(images) > MAX_IMAGES_PER_PRODUCT:
            raise HTTPException(400, f"Max {MAX_IMAGES_PER_PRODUCT} images allowed")

        for image in images:
            if not _is_image_file(image):
                raise HTTPException(400, "Only image files allowed")

            filename = safe_filename(f"{int(time.time())}_{image.filename}")
            image_path = await validate_and_save_image(image, filename)
            image_urls.append(image_path)

    # ---------------- VIDEO UPLOAD ----------------
    video_url: Optional[str] = None
    if video:
        video_url = await save_file_with_limit(
            file=video,
            allowed_types=ALLOWED_VIDEO_TYPES,
            max_size=MAX_VIDEO_SIZE,
            upload_dir="static/uploads/product_videos"  # ✅ Now this works
        )

    # ---------------- PDF UPLOAD ----------------
    pdf_url: Optional[str] = None
    if pdf:
        pdf_url = await save_file_with_limit(
            file=pdf,
            allowed_types=ALLOWED_PDF_TYPES,
            max_size=MAX_PDF_SIZE,
            upload_dir="static/uploads/product_pdfs"  # ✅ Now this works
        )

    # ---------------- PARSE LIST FIELDS ----------------
    sizes_list = safe_parse_list_field(sizes)
    colors_list = safe_parse_list_field(colors)
    highlights_list = safe_parse_list_field(highlights)

    specs_text = specifications.strip() if specifications else None

    # ---------------- CREATE PRODUCT ----------------
    product = Product(
        name=name,
        description=description,
        details=details,
        specifications=specs_text,
        highlights=list_to_db_str(highlights_list),
        price=price,
        discount_price=discount_price,
        category_id=category.id,
        subcategory_id=sub.id if sub else None,
        sizes=list_to_db_str(sizes_list),
        colors=list_to_db_str(colors_list),
        video_url=video_url,
        pdf_url=pdf_url,
        in_stock=in_stock,
        rating=0.0,
        reviews=0,
        featured=featured,
        best_seller=best_seller,
        new_arrival=new_arrival,
        sgst=sgst,
        cgst=cgst,
    )

    product.images = image_urls

    db.add(product)
    db.commit()
    db.refresh(product)

    print(
        "RETURN TYPES:",
        type(product.video_url),
        type(product.pdf_url),
        product.video_url,
        product.pdf_url
    )

    return ProductOut.model_validate(product)




@router.get("/product/{product_id}", response_model=ProductOut)
def get_product(product_id: int, db: Session = Depends(get_db)): # admin=Depends(get_current_admin_user)
    if product_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid product_id")
    p = db.query(Product).filter(Product.id == product_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    return p


@router.get("/products/best-sellers")
def get_best_sellers(
    db: Session = Depends(get_db),

    # 🔹 Pagination
    page: int = Query(1, ge=1),
    limit: int = Query(200, ge=1, le=100),
):
    """List Best Seller products with pagination"""

    # ⬇️ newest first (optional but recommended)
    query = (
        db.query(Product)
        .filter(Product.best_seller == True)
        .order_by(Product.id.desc())
    )

    # 📄 Pagination
    offset = (page - 1) * limit
    products = query.offset(offset).limit(limit).all()

    def to_list(value):
        if isinstance(value, str):
            return [v.strip() for v in value.split(",") if v.strip()]
        elif isinstance(value, list):
            return [v.strip() for v in value if isinstance(v, str) and v.strip()]
        return []

    result = []
    for p in products:
        result.append({
            "id": p.id,
            "name": p.name,
            "description": p.description,
            "details": p.details,
            "specifications": p.specifications,
            "highlights": to_list(p.highlights),
            "price": p.price,
            "discount_price": p.discount_price,
            "category_id": p.category_id,
            "subcategory_id": p.subcategory_id,
            "category": p.category.name if p.category else None,
            "subcategory": p.subcategory.name if p.subcategory else None,
            "sizes": to_list(p.sizes),
            "colors": to_list(p.colors),
            "images": to_list(p.images),
            "video_url": p.video_url,
            "pdf_url": p.pdf_url,
            "in_stock": p.in_stock,
            "is_active": p.is_active,
            "hsn_code": p.hsn_code,
            "sgst": p.sgst,
            "cgst": p.cgst,
            "rating": p.rating,
            "reviews": p.reviews,
            "featured": p.featured,
            "best_seller": p.best_seller,
            "new_arrival": p.new_arrival,
            "created_at": p.created_at,
        })

    return result

def parse_images(value):
    if not value:
        return []

    # If value is JSON list
    try:
        if isinstance(value, str) and value.strip().startswith("["):
            return json.loads(value)
    except:
        pass

    # If comma-separated string
    if isinstance(value, str) and "," in value:
        return [img.strip() for img in value.split(",")]

    # Single string image
    return [value] if isinstance(value, str) else []

@router.post("/products/best-seller")
def update_best_seller_status(payload: BestSellerRequest, db: Session = Depends(get_db)):

    product = db.query(Product).filter(Product.id == payload.product_id).first()

    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    try:
        product.best_seller = payload.is_best_seller
        db.commit()

        return {
            "success": True,
            "product_id": payload.product_id,
            "best_seller": payload.is_best_seller,
            "message": "Best Seller status updated successfully"
        }

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to update: {str(e)}")


@router.get("/products", response_model=List[ProductOut])
def list_products(
        db: Session = Depends(get_db),

        # 🔹 Pagination
        page: int = Query(1, ge=1),
        limit: int = Query(20, ge=1, le=100),

        # 🔹 Search
        search: Optional[str] = Query(None, description="Search by product name or description"),

        # 🔹 Category Filters
        category_id: Optional[int] = Query(None, description="Filter by category ID"),
        subcategory_id: Optional[int] = Query(None, description="Filter by subcategory ID"),

        # 🔹 Price Filters
        min_price: Optional[float] = Query(None, ge=0, description="Minimum price"),
        max_price: Optional[float] = Query(None, ge=0, description="Maximum price"),

        # 🔹 Stock Filter
        in_stock: Optional[bool] = Query(None, description="Filter by stock availability"),

        # 🔹 Product Flags
        featured: Optional[bool] = Query(None, description="Filter featured products"),
        best_seller: Optional[bool] = Query(None, description="Filter best sellers"),
        new_arrival: Optional[bool] = Query(None, description="Filter new arrivals"),

        # 🔹 Rating Filter
        min_rating: Optional[float] = Query(None, ge=0, le=5, description="Minimum rating"),

        # 🔹 Attribute Filters
        sizes: Optional[str] = Query(None, description="Filter by sizes (comma-separated)"),
        colors: Optional[str] = Query(None, description="Filter by colors (comma-separated)"),

        # 🔹 Sorting
        sort_by: Optional[str] = Query("newest",
                                       description="Sort by: newest, oldest, price_low, price_high, rating, popular"),
):
    """
    List products with comprehensive filtering:
    - Pagination
    - Search by name/description
    - Category and subcategory filters
    - Price range filtering
    - Stock availability
    - Product flags (featured, best_seller, new_arrival)
    - Rating filter
    - Size and color filters
    - Multiple sorting options
    """

    query = db.query(Product)

    # 🔍 SEARCH FILTER
    if search:
        term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Product.name.ilike(term),
                Product.description.ilike(term),
            )
        )

    # 🏷️ CATEGORY FILTERS
    if category_id:
        query = query.filter(Product.category_id == category_id)

    if subcategory_id:
        query = query.filter(Product.subcategory_id == subcategory_id)

    # 💰 PRICE RANGE FILTERS
    if min_price is not None:
        # Filter by discount_price if available, otherwise by regular price
        query = query.filter(
            or_(
                and_(Product.discount_price.isnot(None), Product.discount_price >= min_price),
                and_(Product.discount_price.is_(None), Product.price >= min_price)
            )
        )

    if max_price is not None:
        query = query.filter(
            or_(
                and_(Product.discount_price.isnot(None), Product.discount_price <= max_price),
                and_(Product.discount_price.is_(None), Product.price <= max_price)
            )
        )

    # 📦 STOCK FILTER
    if in_stock is not None:
        query = query.filter(Product.in_stock == in_stock)

    # 🏆 PRODUCT FLAGS FILTERS
    if featured is not None:
        query = query.filter(Product.featured == featured)

    if best_seller is not None:
        query = query.filter(Product.best_seller == best_seller)

    if new_arrival is not None:
        query = query.filter(Product.new_arrival == new_arrival)

    # ⭐ RATING FILTER
    if min_rating is not None:
        query = query.filter(Product.rating >= min_rating)

    # 🎨 SIZE FILTER
    if sizes:
        size_list = [s.strip() for s in sizes.split(",") if s.strip()]
        size_filters = [Product.sizes.ilike(f"%{size}%") for size in size_list]
        query = query.filter(or_(*size_filters))

    # 🌈 COLOR FILTER
    if colors:
        color_list = [c.strip() for c in colors.split(",") if c.strip()]
        color_filters = [Product.colors.ilike(f"%{color}%") for color in color_list]
        query = query.filter(or_(*color_filters))

    # ⬆️⬇️ SORTING
    if sort_by == "oldest":
        query = query.order_by(asc(Product.id))
    elif sort_by == "price_low":
        query = query.order_by(
            asc(
                case(
                    (Product.discount_price.isnot(None), Product.discount_price),
                    else_=Product.price
                )
            )
        )
    elif sort_by == "price_high":
        query = query.order_by(
            desc(
                case(
                    (Product.discount_price.isnot(None), Product.discount_price),
                    else_=Product.price
                )
            )
        )
    elif sort_by == "rating":
        query = query.order_by(desc(Product.rating))
    elif sort_by == "popular":
        query = query.order_by(desc(Product.reviews))
    else:  # Default: newest
        query = query.order_by(desc(Product.id))

    # 📄 PAGINATION
    offset = (page - 1) * limit
    products = query.offset(offset).limit(limit).all()

    def to_list(value):
        if isinstance(value, str):
            return [v.strip() for v in value.split(",") if v.strip()]
        elif isinstance(value, list):
            return [v.strip() for v in value if isinstance(v, str) and v.strip()]
        return []

    result = []
    for p in products:
        result.append({
            "id": p.id,
            "name": p.name,
            "description": p.description,
            "details": p.details,
            "specifications": p.specifications,
            "highlights": to_list(p.highlights),
            "price": p.price,
            "discount_price": p.discount_price,
            "category_id": p.category_id,
            "subcategory_id": p.subcategory_id,
            "category": p.category.name if p.category else None,
            "subcategory": p.subcategory.name if p.subcategory else None,
            "sizes": to_list(p.sizes),
            "colors": to_list(p.colors),
            "images": to_list(p.images),
            "video_url": p.video_url,
            "pdf_url": p.pdf_url,
            "in_stock": p.in_stock,
            "is_active": p.is_active,
            "hsn_code": p.hsn_code,
            "sgst": p.sgst,
            "cgst": p.cgst,
            "rating": p.rating,
            "reviews": p.reviews,
            "featured": p.featured,
            "best_seller": p.best_seller,
            "new_arrival": p.new_arrival,
            "created_at": p.created_at,
        })

    return result


@router.put("/product/{product_id}", response_model=ProductOut)
async def update_product(
        product_id: int,
        name: Optional[str] = Form(None),
        description: Optional[str] = Form(None),
        price: Optional[float] = Form(None),
        discount_price: Optional[float] = Form(None),
        category_id: Optional[int] = Form(None),
        subcategory_id: Optional[int] = Form(None),
        rating: Optional[float] = Form(None),
        reviews: Optional[int] = Form(0),
        colors: Optional[str] = Form(None),
        sizes: Optional[str] = Form(None),
        images: Optional[List[UploadFile]] = File(None),
        replace_images_json: Optional[str] = Form(None),

        # ⭐ NEW FILE FIELDS
        video: Optional[UploadFile] = File(None),
        pdf: Optional[UploadFile] = File(None),

        in_stock: Optional[bool] = Form(None),
        featured: Optional[bool] = Form(None),
        best_seller: Optional[bool] = Form(None),
        new_arrival: Optional[bool] = Form(None),
        highlights: Optional[str] = Form(None),
        specifications: Optional[str] = Form(None),
        details: Optional[str] = Form(None),
        sgst: Optional[float] = Form(None),
        cgst: Optional[float] = Form(None),

        db: Session = Depends(get_db),
        admin=Depends(get_current_admin_user)
):
    import os
    import json
    import time

    if product_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid product_id")

    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    # --------------------------------------
    # BASIC FIELD UPDATES
    # --------------------------------------
    if name is not None:
        product.name = validate_name(name)

    if description is not None:
        product.description = description

    if details is not None:
        product.details = details

    if price is not None:
        product.price = validate_price(price)

    if discount_price is not None:
        discount_price = validate_price(discount_price)
        compare_price = price if price is not None else product.price

        if discount_price > compare_price:
            raise HTTPException(status_code=400, detail="discount_price cannot exceed price")

        product.discount_price = discount_price

    if category_id is not None:
        cat = ensure_category_exists(db, category_id)
        product.category_id = cat.id

    if subcategory_id is not None:
        sub = ensure_subcategory_exists(db, subcategory_id)
        product.subcategory_id = sub.id

    if highlights is not None:
        product.highlights = list_to_db_str(safe_parse_list_field(highlights))

    if specifications is not None:
        product.specifications = specifications.strip() if specifications.strip() else None

    # --------------------------------------
    # RATING & REVIEWS
    # --------------------------------------
    if rating is not None:
        try:
            r = float(rating)
        except:
            raise HTTPException(status_code=400, detail="rating must be a number")

        if not (0 <= r <= 5):
            raise HTTPException(status_code=400, detail="rating must be 0–5")

        product.rating = round(r, 1)

    if reviews is not None:
        try:
            rv = int(reviews)
        except:
            raise HTTPException(status_code=400, detail="reviews must be integer")

        if rv < 0:
            raise HTTPException(status_code=400, detail="reviews cannot be negative")

        product.reviews = rv

    # --------------------------------------
    # GST SPLIT
    # --------------------------------------
    if sgst is not None:
        product.sgst = _validate_gst_percent(sgst, "SGST")
    if cgst is not None:
        product.cgst = _validate_gst_percent(cgst, "CGST")

    # --------------------------------------
    # COLORS & SIZES
    # --------------------------------------
    if colors is not None:
        product.colors = list_to_db_str(safe_parse_list_field(colors))

    if sizes is not None:
        product.sizes = list_to_db_str(safe_parse_list_field(sizes))

    # --------------------------------------
    # IMAGE HANDLING
    # --------------------------------------
    replace_images = None

    if replace_images_json:
        try:
            replace_images = json.loads(replace_images_json)
            if not isinstance(replace_images, list):
                raise ValueError
        except:
            raise HTTPException(status_code=400, detail="replace_images_json must be list")

    appended_urls = []

    if images:
        for up in images:
            if not _is_image_file(up):
                raise HTTPException(status_code=400, detail="All uploads must be images.")

            filename = safe_filename(f"{int(time.time())}_{up.filename}")
            saved = await validate_and_save_image(up, filename)
            appended_urls.append(saved)

    if replace_images is not None:
        product.images = list_to_db_str(replace_images + appended_urls)
    elif appended_urls:
        existing = safe_parse_list_field(product.images)
        product.images = list_to_db_str(existing + appended_urls)

    # --------------------------------------
    # ⭐ VIDEO UPLOAD
    # --------------------------------------
    if video:
        if not video.filename.lower().endswith((".mp4", ".mov", ".mkv")):
            raise HTTPException(status_code=400, detail="Video must be MP4, MOV or MKV")

        video_folder = "static/videos"
        os.makedirs(video_folder, exist_ok=True)

        filename = safe_filename(f"video_{product_id}_{int(time.time())}_{video.filename}")
        save_path = f"{video_folder}/{filename}"

        with open(save_path, "wb") as buffer:
            buffer.write(await video.read())

        product.video = save_path

    # --------------------------------------
    # ⭐ PDF UPLOAD
    # --------------------------------------
    if pdf:
        if not pdf.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=400, detail="PDF must be .pdf")

        pdf_folder = "static/pdfs"
        os.makedirs(pdf_folder, exist_ok=True)

        filename = safe_filename(f"pdf_{product_id}_{int(time.time())}_{pdf.filename}")
        save_path = f"{pdf_folder}/{filename}"

        with open(save_path, "wb") as buffer:
            buffer.write(await pdf.read())

        product.pdf = save_path

    # --------------------------------------
    # FLAGS
    # --------------------------------------
    if in_stock is not None:
        product.in_stock = bool(in_stock)

    if featured is not None:
        product.featured = bool(featured)

    if best_seller is not None:
        product.best_seller = bool(best_seller)

    if new_arrival is not None:
        product.new_arrival = bool(new_arrival)

    db.commit()
    db.refresh(product)

    return product




class ProductVisibilityUpdate(BaseModel):
    product_ids: List[int]
    is_active: bool


@router.patch("/products/visibility")
def set_products_visibility(
    payload: ProductVisibilityUpdate,
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Show or hide products on the storefront in one call.

    Hiding flips is_active only — the catalogue entry, images, stock and
    order history stay untouched, and the product vanishes from storefront
    listings and its own page until shown again. Takes one id or a whole
    selection; ids that don't exist are reported back, not errors.
    """
    ids = list({int(i) for i in payload.product_ids})
    if not ids:
        raise HTTPException(status_code=400, detail="No products selected")
    if len(ids) > 200:
        raise HTTPException(status_code=400, detail="Too many products in one request (max 200)")

    products = db.query(Product).filter(Product.id.in_(ids)).all()
    for product in products:
        product.is_active = payload.is_active

    db.commit()

    found_ids = [p.id for p in products]
    missing = sorted(set(ids) - set(found_ids))
    verb = "shown" if payload.is_active else "hidden"
    return {
        "updated": len(found_ids),
        "product_ids": found_ids,
        "missing_ids": missing,
        "is_active": payload.is_active,
        "message": f"{len(found_ids)} product(s) {verb}",
    }


def _remove_or_archive_product(db: Session, product) -> dict:
    """The one true product-removal path, shared by single and bulk delete.

    Carts and wishlists only point at the catalogue — cleared either way.
    (Their FKs have no ON DELETE action; deleting the product first raised
    1451.) Order and invoice lines are records of money that changed hands
    and are never deleted: a product that appears in them is archived —
    hidden from the storefront, kept so past orders stay whole. Only a
    never-sold product is removed outright (reviews cascade, reels keep
    playing with their product link cleared). Does not commit.
    """
    from app.models.cart import CartItem
    from app.models.invoice import InvoiceItem

    db.query(CartItem).filter(
        CartItem.product_id == product.id
    ).delete(synchronize_session=False)
    db.query(WishlistItem).filter(
        WishlistItem.product_id == product.id
    ).delete(synchronize_session=False)

    ordered = db.query(OrderItem).filter(OrderItem.product_id == product.id).count()
    invoiced = db.query(InvoiceItem).filter(InvoiceItem.product_id == product.id).count()

    if ordered or invoiced:
        product.is_active = False
        product.in_stock = False
        return {"archived": True, "ordered": ordered, "invoiced": invoiced, "name": product.name}

    db.delete(product)
    return {"archived": False, "name": product.name}


@router.delete("/product/{product_id}", status_code=200)
def delete_product(
    product_id: int,
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)
):
    if product_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid product_id")

    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    outcome = _remove_or_archive_product(db, product)
    db.commit()

    if outcome["archived"]:
        return {
            "status": True,
            "archived": True,
            "message": (
                f"'{outcome['name']}' appears in {outcome['ordered']} past order line(s)"
                + (f" and {outcome['invoiced']} invoice line(s)" if outcome["invoiced"] else "")
                + " — order history can't lose it, so it was archived instead: "
                  "hidden from the store and removed from carts and wishlists."
            ),
        }

    return {
        "status": True,
        "archived": False,
        "message": "Product deleted permanently"
    }


class ProductBulkDelete(BaseModel):
    product_ids: List[int]


@router.post("/products/bulk-delete")
def bulk_delete_products(
    payload: ProductBulkDelete,
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Delete a selection of products under the same rules as single delete:
    never-sold ones go permanently, ones in past orders are archived."""
    ids = list({int(i) for i in payload.product_ids})
    if not ids:
        raise HTTPException(status_code=400, detail="No products selected")
    if len(ids) > 200:
        raise HTTPException(status_code=400, detail="Too many products in one request (max 200)")

    products = db.query(Product).filter(Product.id.in_(ids)).all()
    deleted = archived = 0
    for product in products:
        if _remove_or_archive_product(db, product)["archived"]:
            archived += 1
        else:
            deleted += 1
    db.commit()

    missing = sorted(set(ids) - {p.id for p in products})
    parts = []
    if deleted:
        parts.append(f"{deleted} deleted permanently")
    if archived:
        parts.append(f"{archived} archived (in past orders — hidden from the store instead)")
    return {
        "status": True,
        "deleted": deleted,
        "archived": archived,
        "missing_ids": missing,
        "message": "; ".join(parts) or "No matching products found",
    }


@router.get("/products/export")
def export_products(
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """The whole catalogue as a spreadsheet — .xlsx via openpyxl, or a CSV
    that Excel opens just the same if openpyxl isn't installed."""
    from datetime import date
    from io import BytesIO, StringIO
    from fastapi.responses import StreamingResponse
    from sqlalchemy.orm import joinedload

    products = (
        db.query(Product)
        .options(joinedload(Product.category), joinedload(Product.subcategory))
        .order_by(Product.id)
        .all()
    )

    headers = [
        "ID", "Name", "Description", "Category", "Subcategory",
        "Price", "Discount Price", "Sizes", "Colors", "HSN Code",
        "SGST %", "CGST %",
        "In Stock", "Visible", "Featured", "Best Seller", "New Arrival",
        "Rating", "Reviews", "Created At", "Images",
    ]

    def cell_list(value):
        if isinstance(value, list):
            return ", ".join(str(v) for v in value)
        return value or ""

    def row(p):
        return [
            p.id,
            p.name,
            p.description or "",
            p.category.name if p.category else "",
            p.subcategory.name if p.subcategory else "",
            p.price,
            p.discount_price or "",
            cell_list(p.sizes),
            cell_list(p.colors),
            p.hsn_code or "",
            p.sgst if p.sgst is not None else "",
            p.cgst if p.cgst is not None else "",
            "Yes" if p.in_stock else "No",
            "No" if p.is_active is False else "Yes",
            "Yes" if p.featured else "No",
            "Yes" if p.best_seller else "No",
            "Yes" if p.new_arrival else "No",
            p.rating or 0,
            p.reviews or 0,
            p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else "",
            cell_list(p.images),
        ]

    stamp = date.today().isoformat()
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(headers)
        for p in products:
            ws.append(row(p))
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="vastraliya-products-{stamp}.xlsx"'},
        )
    except ImportError:
        import csv

        text = StringIO()
        writer = csv.writer(text)
        writer.writerow(headers)
        for p in products:
            writer.writerow(row(p))
        return StreamingResponse(
            BytesIO(text.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="vastraliya-products-{stamp}.csv"'},
        )


@router.post("/products/import")
async def import_products(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Bulk add or edit products from the Download Excel layout.

    Matched by header name, order free. A row WITH an ID updates that
    product — only non-empty cells overwrite. A row WITHOUT one creates a
    product and needs Name, Category (by name) and Price. Sizes/Colors are
    comma-separated; In Stock/Visible/Featured/Best Seller/New Arrival take
    Yes or No. Rating, Reviews, Created At and Images columns are ignored —
    ratings are computed from reviews, and photos are added in the editor.
    Bad rows are reported individually and do not stop the rest.
    """
    from io import BytesIO
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail=(
                "Excel support is not installed on the server. Run: "
                ".venv/bin/pip install openpyxl — then restart the service."
            ),
        )

    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 5 MB)")
    try:
        sheet = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Could not read the file — upload the .xlsx sheet from Download Excel",
        )

    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    except StopIteration:
        raise HTTPException(status_code=400, detail="The sheet is empty")

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    idx = {
        "id": col("id", "product id"),
        "name": col("name"),
        "description": col("description"),
        "category": col("category"),
        "subcategory": col("subcategory"),
        "price": col("price"),
        "discount": col("discount price"),
        "sizes": col("sizes"),
        "colors": col("colors"),
        "hsn": col("hsn code"),
        "sgst": col("sgst %", "sgst"),
        "cgst": col("cgst %", "cgst"),
        "in_stock": col("in stock"),
        "visible": col("visible"),
        "featured": col("featured"),
        "best_seller": col("best seller"),
        "new_arrival": col("new arrival"),
    }
    if idx["name"] is None or idx["price"] is None or idx["category"] is None:
        raise HTTPException(
            status_code=400,
            detail="Missing columns — the sheet needs at least Name, Category and Price",
        )

    categories = {c.name.strip().lower(): c for c in db.query(models.Category).all()}
    subcats = {}
    for s in db.query(models.SubCategory).all():
        subcats[(s.category_id, (s.name or "").strip().lower())] = s

    def cell(row, key):
        i = idx[key]
        if i is None or i >= len(row) or row[i] is None:
            return None
        v = str(row[i]).strip()
        return v or None

    def as_bool(v):
        if v is None:
            return None
        return v.strip().lower() in ("yes", "true", "1", "y")

    def as_float(v):
        try:
            return float(v) if v is not None else None
        except ValueError:
            return None

    created = updated = 0
    errors = []

    for line_no, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        if line_no > 2001:
            errors.append({"row": line_no, "error": "Too many rows (max 2000 per upload)"})
            break

        try:
            pid_raw = cell(row, "id")
            product = None
            if pid_raw:
                try:
                    pid = int(float(pid_raw))
                except ValueError:
                    errors.append({"row": line_no, "error": f"Bad ID '{pid_raw}'"})
                    continue
                product = db.query(Product).filter(Product.id == pid).first()
                if not product:
                    errors.append({"row": line_no, "error": f"Product ID {pid} not found"})
                    continue

            name = cell(row, "name")
            price = as_float(cell(row, "price"))
            cat_name = cell(row, "category")
            category = categories.get(cat_name.lower()) if cat_name else None

            if product is None:
                if not name:
                    errors.append({"row": line_no, "error": "Name is required for a new product"})
                    continue
                if price is None or price <= 0:
                    errors.append({"row": line_no, "error": "Price must be a number above 0"})
                    continue
                if category is None:
                    errors.append({"row": line_no, "error": f"Category '{cat_name}' not found"})
                    continue
                product = Product(
                    name=name, price=price, category_id=category.id,
                    in_stock=True, is_active=True, rating=0.0, reviews=0,
                )
                db.add(product)
                created += 1
            else:
                if name:
                    product.name = name
                if price is not None:
                    if price <= 0:
                        errors.append({"row": line_no, "error": "Price must be above 0"})
                        continue
                    product.price = price
                if cat_name:
                    if category is None:
                        errors.append({"row": line_no, "error": f"Category '{cat_name}' not found"})
                        continue
                    product.category_id = category.id
                updated += 1

            sub_name = cell(row, "subcategory")
            if sub_name:
                sub = subcats.get((product.category_id, sub_name.lower()))
                if sub:
                    product.subcategory_id = sub.id
                else:
                    errors.append({"row": line_no, "error": f"Subcategory '{sub_name}' not in that category (other fields saved)"})

            desc = cell(row, "description")
            if desc is not None:
                product.description = desc

            discount = as_float(cell(row, "discount"))
            if discount is not None:
                base = product.price or 0
                if discount <= 0:
                    product.discount_price = None
                elif discount >= base:
                    errors.append({"row": line_no, "error": "Discount Price must be below Price (other fields saved)"})
                else:
                    product.discount_price = discount

            sizes = cell(row, "sizes")
            if sizes is not None:
                product.sizes = ",".join(p.strip() for p in sizes.split(",") if p.strip())
            colors = cell(row, "colors")
            if colors is not None:
                product.colors = ",".join(p.strip() for p in colors.split(",") if p.strip())

            hsn = cell(row, "hsn")
            if hsn is not None:
                product.hsn_code = hsn[:20]
            sgst = as_float(cell(row, "sgst"))
            if sgst is not None:
                product.sgst = _validate_gst_percent(sgst, "SGST")
            cgst = as_float(cell(row, "cgst"))
            if cgst is not None:
                product.cgst = _validate_gst_percent(cgst, "CGST")

            for key, attr in (
                ("in_stock", "in_stock"), ("visible", "is_active"),
                ("featured", "featured"), ("best_seller", "best_seller"),
                ("new_arrival", "new_arrival"),
            ):
                b = as_bool(cell(row, key))
                if b is not None:
                    setattr(product, attr, b)
        except HTTPException as exc:
            errors.append({"row": line_no, "error": str(exc.detail)})
            continue

    db.commit()

    parts = []
    if created:
        parts.append(f"{created} product(s) added")
    if updated:
        parts.append(f"{updated} updated")
    if errors:
        parts.append(f"{len(errors)} problem(s) reported")
    return {
        "status": True,
        "created": created,
        "updated": updated,
        "errors": errors[:50],
        "message": ", ".join(parts) or "Nothing imported",
    }


# ------------------- REVIEWS: EXCEL EXPORT / IMPORT -------------------

REVIEW_SHEET_HEADERS = [
    "Review ID", "Product ID", "Product Name", "Customer Name", "Rating", "Comment", "Created At",
]


def _recalc_review_aggregates(db: Session, product_ids: set) -> None:
    """Keep Product.rating (average) and Product.reviews (count) truthful."""
    from app.models import ProductReview

    # The session may not autoflush — the aggregates must see rows added
    # and edited in this request.
    db.flush()

    for pid in product_ids:
        product = db.query(Product).filter(Product.id == pid).first()
        if not product:
            continue
        avg = db.query(func.avg(ProductReview.rating)).filter(ProductReview.product_id == pid).scalar()
        cnt = db.query(func.count(ProductReview.id)).filter(ProductReview.product_id == pid).scalar()
        product.rating = round(float(avg), 1) if avg is not None else 0.0
        product.reviews = int(cnt or 0)


@router.get("/reviews/export")
def export_reviews(
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Every review as a spreadsheet, in the same layout the import reads —
    download, edit, re-upload and rows with a Review ID update in place."""
    from datetime import date
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from sqlalchemy.orm import joinedload
    from app.models import ProductReview

    reviews = (
        db.query(ProductReview)
        .options(joinedload(ProductReview.product))
        .order_by(ProductReview.id)
        .all()
    )

    def row(r):
        return [
            r.id,
            r.product_id,
            r.product.name if r.product else "",
            r.name,
            r.rating,
            r.comment or "",
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        ]

    stamp = date.today().isoformat()
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Reviews"
        ws.append(REVIEW_SHEET_HEADERS)
        for r in reviews:
            ws.append(row(r))
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="vastraliya-reviews-{stamp}.xlsx"'},
        )
    except ImportError:
        # No openpyxl on this install — a CSV opens in Excel just the same.
        import csv
        from io import StringIO

        text = StringIO()
        writer = csv.writer(text)
        writer.writerow(REVIEW_SHEET_HEADERS)
        for r in reviews:
            writer.writerow(row(r))
        return StreamingResponse(
            BytesIO(text.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="vastraliya-reviews-{stamp}.csv"'},
        )


@router.post("/reviews/import")
async def import_reviews(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    """Bulk add or edit reviews from a spreadsheet.

    Columns (matched by header name, order free): Review ID, Product ID,
    Customer Name, Rating, Comment. A row WITH a Review ID updates that
    review's name/rating/comment; a row WITHOUT one creates a new review
    on the given Product ID. Bad rows are reported back individually and
    do not stop the rest. Imported reviews belong to the importing admin
    account; the storefront shows the Customer Name column.
    """
    from io import BytesIO
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail=(
                "Excel support is not installed on the server. Run: "
                ".venv/bin/pip install openpyxl — then restart the service."
            ),
        )
    from app.models import ProductReview

    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 5 MB)")
    try:
        ws = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Could not read the file — upload the .xlsx sheet from Download Excel",
        )

    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    except StopIteration:
        raise HTTPException(status_code=400, detail="The sheet is empty")

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_id = col("review id", "id")
    i_pid = col("product id")
    i_name = col("customer name", "name")
    i_rating = col("rating")
    i_comment = col("comment", "review")

    if i_pid is None or i_name is None or i_rating is None:
        raise HTTPException(
            status_code=400,
            detail="Missing columns — the sheet needs Product ID, Customer Name and Rating",
        )

    def cell(row, i):
        if i is None or i >= len(row) or row[i] is None:
            return None
        v = str(row[i]).strip()
        return v or None

    created = updated = 0
    errors = []
    touched_products: set = set()

    for line_no, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        if line_no > 2001:
            errors.append({"row": line_no, "error": "Too many rows (max 2000 per upload)"})
            break

        name = cell(row, i_name)
        rating_raw = cell(row, i_rating)
        comment = cell(row, i_comment)
        try:
            rating = float(rating_raw) if rating_raw is not None else None
        except ValueError:
            rating = None
        if rating is None or not (1 <= rating <= 5):
            errors.append({"row": line_no, "error": "Rating must be a number from 1 to 5"})
            continue
        if not name:
            errors.append({"row": line_no, "error": "Customer Name is required"})
            continue

        review_id = cell(row, i_id)
        if review_id:
            try:
                rid = int(float(review_id))
            except ValueError:
                errors.append({"row": line_no, "error": f"Bad Review ID '{review_id}'"})
                continue
            review = db.query(ProductReview).filter(ProductReview.id == rid).first()
            if not review:
                errors.append({"row": line_no, "error": f"Review ID {rid} not found"})
                continue
            review.name = name[:120]
            review.rating = rating
            review.comment = comment
            touched_products.add(review.product_id)
            updated += 1
        else:
            pid_raw = cell(row, i_pid)
            try:
                pid = int(float(pid_raw)) if pid_raw is not None else None
            except ValueError:
                pid = None
            if pid is None or not db.query(Product.id).filter(Product.id == pid).first():
                errors.append({"row": line_no, "error": f"Product ID '{pid_raw}' not found"})
                continue
            db.add(ProductReview(
                product_id=pid,
                user_id=admin.id,
                name=name[:120],
                rating=rating,
                comment=comment,
            ))
            touched_products.add(pid)
            created += 1

    _recalc_review_aggregates(db, touched_products)
    db.commit()

    parts = []
    if created:
        parts.append(f"{created} review(s) added")
    if updated:
        parts.append(f"{updated} updated")
    if errors:
        parts.append(f"{len(errors)} row(s) skipped")
    return {
        "status": True,
        "created": created,
        "updated": updated,
        "errors": errors[:50],
        "message": ", ".join(parts) or "Nothing imported",
    }






def _recalculate_product_rating(db: Session, product: Product):
    rows = db.query(models.ProductReview).filter(models.ProductReview.product_id == product.id).all()
    total = len(rows)
    if total == 0:
        product.rating = 0.0
        product.reviews = 0
    else:
        s = sum((r.rating or 0) for r in rows)
        product.rating = round(s / total, 2)
        product.reviews = total
    db.commit()
    db.refresh(product)



# -------------------------------------------------------------------
# 📊 ADMIN ANALYTICS
# -------------------------------------------------------------------

@router.get("/analytics/monthly-orders")
def get_monthly_orders(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user),
):
    """Number of orders per month."""
    year_expr = extract('year', models.Order.created_at)
    month_expr = extract('month', models.Order.created_at)

    result = (
        db.query(
            year_expr.label("year"),
            month_expr.label("month"),
            func.count(models.Order.id).label("total_orders"),
        )
        .select_from(models.Order)
        .group_by(year_expr, month_expr)
        .order_by(year_expr, month_expr)
        .all()
    )

    # If using SQLite helpers above, cast to int when building the response:
    return [
        {"year": int(r.year), "month": int(r.month), "orders": r.total_orders}
        for r in result
    ]


@router.get("/analytics/daily-orders")
def get_daily_orders(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user),
):
    """Orders created per day for last 7 days."""
    seven_days_ago = datetime.utcnow() - timedelta(days=7)

    # For SQLite, replace day_expr with the commented version at top
    day_expr = func.date(models.Order.created_at)

    result = (
        db.query(
            day_expr.label("day"),
            func.count(models.Order.id).label("total_orders"),
        )
        .select_from(models.Order)
        .filter(models.Order.created_at >= seven_days_ago)
        .group_by(day_expr)
        .order_by(day_expr)
        .all()
    )
    return [{"day": str(r.day), "orders": r.total_orders} for r in result]


@router.get("/analytics/top-customers")
def get_top_customers(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user),
):
    """Top 10 customers by number of orders."""
    # Portable name expression (works on SQLite/Postgres/MySQL)
    name_expr = (func.coalesce(User.first_name, '') + literal(' ') +
                 func.coalesce(User.last_name, ''))

    result = (
        db.query(
            User.id.label("user_id"),
            name_expr.label("name"),
            func.count(models.Order.id).label("order_count"),
        )
        .select_from(models.Order)
        .join(User, models.Order.user_id == User.id)
        .group_by(User.id, User.first_name, User.last_name)
        .order_by(func.count(models.Order.id).desc())
        .limit(10)
        .all()
    )
    return [{"id": r.user_id, "name": r.name, "orders": r.order_count} for r in result]

# -------------------------------------------------------------------
#                        COUPON MANAGEMENT
# -------------------------------------------------------------------

@router.post("/create")
def create_coupon(
    code: str = Form(...),
    description: Optional[str] = Form(None),
    discount_type: str = Form(...),  # "percentage" or "fixed"
    discount_value: float = Form(...),
    minimum_purchase: float = Form(0.0),
    valid_from: datetime = Form(...),
    valid_to: datetime = Form(...),
    max_uses: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user)
):
    """Create a new coupon with validation."""
    if not code or not code.strip():
        raise HTTPException(status_code=400, detail="Coupon code is required.")
    code_clean = code.strip().upper()
    if len(code_clean) > 50:
        raise HTTPException(status_code=400, detail="Coupon code too long (max 50 chars).")

    if discount_type not in ["percentage", "fixed"]:
        raise HTTPException(status_code=400, detail="Invalid discount_type. Must be 'percentage' or 'fixed'.")

    discount_value = validate_price(discount_value)
    if discount_type == "percentage" and discount_value > 100:
        raise HTTPException(status_code=400, detail="percentage discount_value cannot exceed 100.")

    minimum_purchase = validate_price(minimum_purchase)

    if valid_to <= valid_from:
        raise HTTPException(status_code=400, detail="valid_to must be after valid_from.")

    if max_uses is not None:
        try:
            max_uses = int(max_uses)
            if max_uses < 0:
                raise HTTPException(status_code=400, detail="max_uses cannot be negative.")
        except Exception:
            raise HTTPException(status_code=400, detail="max_uses must be an integer.")

    existing = db.query(Coupon).filter(func.lower(Coupon.code) == code_clean.lower()).first()
    if existing:
        raise HTTPException(status_code=400, detail="Coupon code already exists.")

    coupon = Coupon(
        code=code_clean,
        description=(description or "").strip(),
        discount_type=discount_type,
        discount_value=discount_value,
        minimum_purchase=minimum_purchase,
        valid_from=valid_from,
        valid_to=valid_to,
        max_uses=max_uses,
        active=True
    )
    db.add(coupon)
    db.commit()
    db.refresh(coupon)
    return {"message": "Coupon created successfully", "coupon_id": coupon.id}


@router.get("/list")
def list_coupons(

        db: Session = Depends(get_db),
    current_user = Depends(get_current_admin_user)


): # current_user=Depends(get_current_admin_user)
    """List all coupons (auto-deactivates expired ones)."""
    now = datetime.utcnow()

    # Auto-disable expired coupons
    expired = (
        db.query(Coupon)
        .filter(Coupon.valid_to < now, Coupon.active == True)
        .update({"active": False})
    )
    if expired:
        db.commit()

    coupons = db.query(Coupon).order_by(Coupon.created_at.desc()).all() if hasattr(Coupon, "created_at") else db.query(Coupon).all()

    return [
        {
            "id": c.id,
            "code": c.code,
            "description": c.description,
            "discount_type": c.discount_type,
            "discount_value": c.discount_value,
            "minimum_purchase": c.minimum_purchase,
            "valid_from": c.valid_from,
            "valid_to": c.valid_to,
            "max_uses": c.max_uses,
            "used_count": getattr(c, "used_count", 0),
            "active": c.active,
        }
        for c in coupons
    ]


@router.put("/update/{coupon_id}")
def update_coupon(
    coupon_id: int,
    code: Optional[str] = Form(None),
    description: Optional[str] = Form(None),
    discount_type: Optional[str] = Form(None),
    discount_value: Optional[float] = Form(None),
    minimum_purchase: Optional[float] = Form(None),
    valid_from: Optional[datetime] = Form(None),
    valid_to: Optional[datetime] = Form(None),
    max_uses: Optional[int] = Form(None),
    active: Optional[bool] = Form(None),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user)
):
    """Update coupon details."""
    if coupon_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid coupon_id")
    coupon = db.query(Coupon).filter(Coupon.id == coupon_id).first()
    if not coupon:
        raise HTTPException(status_code=404, detail="Coupon not found")

    if discount_type and discount_type not in ["percentage", "fixed"]:
        raise HTTPException(status_code=400, detail="Invalid discount_type")

    if valid_from and valid_to and valid_to <= valid_from:
        raise HTTPException(status_code=400, detail="valid_to must be after valid_from.")

    if code:
        code_clean = code.strip().upper()
        if len(code_clean) > 50:
            raise HTTPException(status_code=400, detail="Coupon code too long (max 50 chars).")
        existing = db.query(Coupon).filter(func.lower(Coupon.code) == code_clean.lower(), Coupon.id != coupon_id).first()
        if existing:
            raise HTTPException(status_code=400, detail="Coupon code already in use.")
        coupon.code = code_clean

    if discount_value is not None:
        discount_value = validate_price(discount_value)
        # if discount_type provided use that, else use coupon.discount_type
        d_type = discount_type if discount_type is not None else coupon.discount_type
        if d_type == "percentage" and discount_value > 100:
            raise HTTPException(status_code=400, detail="percentage discount_value cannot exceed 100.")
        coupon.discount_value = discount_value

    if minimum_purchase is not None:
        coupon.minimum_purchase = validate_price(minimum_purchase)

    if max_uses is not None:
        try:
            max_uses = int(max_uses)
            if max_uses < 0:
                raise HTTPException(status_code=400, detail="max_uses cannot be negative.")
            coupon.max_uses = max_uses
        except Exception:
            raise HTTPException(status_code=400, detail="max_uses must be an integer.")

    updates = {
        "description": description,
        "discount_type": discount_type,
        "valid_from": valid_from,
        "valid_to": valid_to,
        "active": active,
    }

    for field, value in updates.items():
        if value is not None:
            setattr(coupon, field, value)

    db.commit()
    db.refresh(coupon)
    return {"message": "Coupon updated successfully", "coupon": coupon.code}









@router.delete("/delete/{coupon_id}")
def delete_coupon(
    coupon_id: int,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_admin_user)
):
    """Delete a coupon."""
    if coupon_id <= 0:
        raise HTTPException(status_code=400, detail="Invalid coupon_id")
    coupon = db.query(Coupon).filter(Coupon.id == coupon_id).first()
    if not coupon:
        raise HTTPException(status_code=404, detail="Coupon not found")

    db.delete(coupon)
    db.commit()
    return {"message": f"Coupon '{coupon.code}' deleted successfully"}







@router.put("/user/{user_id}/activate", dependencies=[Depends(get_current_admin_user)])
def admin_activate_user(user_id: int, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if not user.deleted_at:
        raise HTTPException(status_code=400, detail="User is already active")

    user.deleted_at = None
    user.active = True
    user.active_at = datetime.utcnow()   # ✅ THIS WAS MISSING

    db.commit()
    db.refresh(user)

    return {"message": f"User {user_id} activated successfully"}


class AdminUserUpdate(BaseModel):
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    role: Optional[str] = None
    address_line1: Optional[str] = None
    address_line2: Optional[str] = None
    city: Optional[str] = None
    region: Optional[str] = None
    postal_code: Optional[str] = None
    country: Optional[str] = None


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

def _count_admins(db: Session) -> int:
    return db.query(User).filter(User.role == "admin", User.deleted_at.is_(None)).count()


@router.put("/user/{user_id}", dependencies=[Depends(get_current_admin_user)])
def admin_update_user(
    user_id: int,
    payload: AdminUserUpdate,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    """Edit a customer's details from the admin dashboard."""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    data = payload.dict(exclude_unset=True)

    if "email" in data:
        email = (data["email"] or "").strip().lower()
        if not EMAIL_RE.match(email):
            raise HTTPException(status_code=400, detail="Enter a valid email address")
        clash = (
            db.query(User)
            .filter(func.lower(User.email) == email, User.id != user_id)
            .first()
        )
        if clash:
            raise HTTPException(status_code=409, detail="Another account already uses that email")
        user.email = email

    if "role" in data:
        role = (data["role"] or "").strip().lower()
        if role not in {"user", "admin"}:
            raise HTTPException(status_code=400, detail="Role must be 'user' or 'admin'")
        # Don't let an admin lock everyone out of the dashboard.
        if user.role == "admin" and role != "admin":
            if user.id == admin_user.id:
                raise HTTPException(status_code=400, detail="You can't remove your own admin access")
            if _count_admins(db) <= 1:
                raise HTTPException(status_code=400, detail="This is the only admin — promote someone else first")
        user.role = role

    for field in (
        "first_name", "last_name", "phone",
        "address_line1", "address_line2", "city", "region", "postal_code", "country",
    ):
        if field in data:
            value = (data[field] or "").strip()
            setattr(user, field, value or None)

    db.commit()
    db.refresh(user)
    return {
        "message": "User updated",
        "user": {
            "id": user.id,
            "first_name": user.first_name,
            "last_name": user.last_name,
            "email": user.email,
            "phone": user.phone,
            "role": user.role,
            "address_line1": user.address_line1,
            "address_line2": user.address_line2,
            "city": user.city,
            "region": user.region,
            "postal_code": user.postal_code,
            "country": user.country,
        },
    }


class AdminSetPassword(BaseModel):
    new_password: str


@router.put("/user/{user_id}/password", dependencies=[Depends(get_current_admin_user)])
def admin_set_user_password(
    user_id: int,
    payload: AdminSetPassword,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    """Set a customer's password from the admin dashboard.

    For the case where a customer can't get back in and the emailed reset link
    isn't reaching them. It deliberately does NOT ask for their old password —
    an admin doesn't have it — so it is a genuine override, and the guards below
    are what stop it becoming an account-takeover tool.
    """
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Your own password still needs your old one. Without that, anyone who
    # borrowed a signed-in dashboard could lock the real admin out of their
    # own account for good; /users/change-password is the door for that.
    if user.id == admin_user.id:
        raise HTTPException(
            status_code=400,
            detail="Change your own password from your profile — it asks for your current one.",
        )

    new_password = payload.new_password or ""
    try:
        validate_password(new_password)
    except PasswordPolicyError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    had_password = bool(user.password)
    user.password = get_password_hash(new_password)

    # Any reset link already in their inbox would still work against the old
    # flow and hand the account back to whoever holds it. Burn them.
    tokens_killed = (
        db.query(PasswordResetToken)
        .filter(PasswordResetToken.email == user.email, PasswordResetToken.used.is_(False))
        .update({PasswordResetToken.used: True}, synchronize_session=False)
    )

    db.commit()

    # Never log the password itself — just enough to answer "who did this".
    logger.info(
        "Admin %s set the password for user %s (reset links invalidated: %s)",
        admin_user.id, user.id, tokens_killed,
    )

    return {
        "message": "Password updated",
        "user_id": user.id,
        "email": user.email,
        # A Google-only account had no password until now; the dashboard says so.
        "first_password": not had_password,
        "reset_links_invalidated": tokens_killed,
    }


@router.delete("/user/{user_id}", dependencies=[Depends(get_current_admin_user)])
def admin_delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    """Permanently delete a customer.

    Refused for anyone with orders or invoices — that history is accounting
    data and outlives the account. Deactivate those instead, which blocks
    login while keeping the record.
    """
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if user.id == admin_user.id:
        raise HTTPException(status_code=400, detail="You can't delete your own account")
    if user.role == "admin" and _count_admins(db) <= 1:
        raise HTTPException(status_code=400, detail="This is the only admin — promote someone else first")

    order_count = db.query(Order).filter(Order.user_id == user_id).count()
    if order_count:
        raise HTTPException(
            status_code=409,
            detail=(
                f"This customer has {order_count} order(s). Deleting them would take the "
                "order history with it — deactivate the account instead."
            ),
        )

    from app.models.invoice import Invoice

    invoice_count = db.query(Invoice).filter(Invoice.user_id == user_id).count()
    if invoice_count:
        raise HTTPException(
            status_code=409,
            detail=(
                f"This customer has {invoice_count} invoice(s). Deactivate the account "
                "instead so the records are kept."
            ),
        )

    # Clear the rows that would otherwise trip a foreign key.
    from app.models.address import Address
    from app.models.cart import CartItem as _CartItem
    from app.models.product import ProductReview, ProductReviewReply
    from app.models.user import SupportMessage as _SupportMessage, UserSettings, PasswordResetToken
    from app.models.wishlist import WishlistItem as _WishlistItem

    for model in (
        Address, UserSettings, _CartItem, _WishlistItem, _SupportMessage, ProductReview,
    ):
        try:
            db.query(model).filter(model.user_id == user_id).delete(synchronize_session=False)
        except Exception:
            db.rollback()
            raise HTTPException(status_code=500, detail=f"Could not clear {model.__name__} rows")
    db.query(ProductReviewReply).filter(ProductReviewReply.admin_id == user_id).delete(
        synchronize_session=False
    )
    # This one is keyed by email rather than user_id.
    db.query(PasswordResetToken).filter(PasswordResetToken.email == user.email).delete(
        synchronize_session=False
    )

    email = user.email
    db.delete(user)
    db.commit()
    return {"message": f"Deleted {email}", "id": user_id}


@router.put("/user/{user_id}/deactivate", dependencies=[Depends(get_current_admin_user)])
def admin_deactivate_user(user_id: int, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if user.deleted_at:
        raise HTTPException(status_code=400, detail="User already deactivated")

    user.deleted_at = datetime.utcnow()
    user.active = False
    user.active_at = None   # ✅ clear timestamp when disabling

    db.commit()
    db.refresh(user)

    return {"message": f"User {user_id} deactivated successfully"}




@router.get("/orders/{order_id}")
def get_order_by_id(
    order_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(get_current_user)
):
    if admin.role != "admin":
        raise HTTPException(status_code=403, detail="Admins only")

    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    return order

@router.put("/orders/{order_id}/status")
def update_order_status(
    order_id: int,
    status: str,
    db: Session = Depends(get_db),
    admin: User = Depends(get_current_user)
):
    if admin.role != "admin":
        raise HTTPException(status_code=403, detail="Admins only")

    status = status.lower().strip()

    if status not in ALLOWED_ORDER_STATUSES:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid status. Allowed: {', '.join(ALLOWED_ORDER_STATUSES)}"
        )

    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    order.status = status
    db.commit()

    db.add(OrderTracking(
        order_id=order.id,
        status=status,
        message=f"Order status updated to {status} by admin"
    ))
    db.commit()

    return {
        "message": "Order status updated",
        "order_id": order.id,
        "status": status
    }

@router.delete("/orders/{order_id}")
def delete_order(
    order_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(get_current_user)
):
    # Only admin allowed
    if admin.role != "admin":
        raise HTTPException(status_code=403, detail="Admins only")

    # Check order exists
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    # Delete order tracking first (if exists)
    db.query(OrderTracking).filter(OrderTracking.order_id == order.id).delete()

    # Delete order
    db.delete(order)
    db.commit()

    return {
        "message": "Order deleted successfully",
        "order_id": order_id
    }

@router.put("/orders/{order_id}/payment-status")
def update_payment_status(
    order_id: int,
    payment_status: str,
    db: Session = Depends(get_db),
    admin: User = Depends(get_current_user)
):
    # Only admin allowed
    if admin.role != "admin":
        raise HTTPException(status_code=403, detail="Admins only")

    payment_status = payment_status.lower().strip()

    ALLOWED_PAYMENT_STATUS = {"paid", "unpaid"}

    if payment_status not in ALLOWED_PAYMENT_STATUS:
        raise HTTPException(
            status_code=400,
            detail="Invalid payment status. Allowed values: paid, unpaid"
        )

    # Find order
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    # Update payment status
    order.payment_status = payment_status
    db.commit()

    return {
        "message": f"Payment status updated to {payment_status}",
        "order_id": order.id,
        "payment_status": payment_status
    }
# Allowed statuses for orders (adjust as your system defines)
ALLOWED_ORDER_STATUSES = {"pending", "processing", "shipped", "delivered", "cancelled", "refunded"}

# Pagination defaults / limits
DEFAULT_PAGE = 1
DEFAULT_PAGE_SIZE = 20
MAX_PAGE_SIZE = 200

# -------------------------
# Request models
# -------------------------
class UpdateOrderStatusRequest(BaseModel):
    status: str

class ReplySupportRequest(BaseModel):
    reply: constr(strip_whitespace=True, min_length=1, max_length=2000)

# -------------------------
# Helpers
# -------------------------
def validate_page_params(page: int, page_size: int):
    if page < 1:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="page must be >= 1")
    if page_size < 1 or page_size > MAX_PAGE_SIZE:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"page_size must be between 1 and {MAX_PAGE_SIZE}")

# -----------------------------------------------------
# 📊 DASHBOARD METRICS
# -----------------------------------------------------
@router.get("/admin_stats")
def get_dashboard_stats(db: Session = Depends(get_db), admin=Depends(get_current_admin_user)):
    """Return key business metrics for the admin dashboard"""
    total_users = db.query(func.count(User.id)).scalar() or 0
    total_orders = db.query(func.count(Order.id)).scalar() or 0
    # coalesce sum to 0.0 when there are no orders or nulls
    total_revenue = db.query(func.coalesce(func.sum(Order.total_amount), 0.0)).scalar() or 0.0
    avg_order_value = (total_revenue / total_orders) if total_orders > 0 else 0.0

    return {
        "total_users": int(total_users),
        "total_orders": int(total_orders),
        "total_revenue": round(float(total_revenue), 2),
        "average_order_value": round(float(avg_order_value), 2)
    }

# -----------------------------------------------------
# 🧾 ORDER MANAGEMENT
# -----------------------------------------------------
DEFAULT_PAGE = 1
DEFAULT_PAGE_SIZE = 10
MAX_PAGE_SIZE = 100


GST_RATE = 0.18

@router.get("/admin_orders")
def list_all_orders(
    page: int = Query(DEFAULT_PAGE, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),

    order_id: Optional[int] = Query(None),
    user_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    payment_status: Optional[str] = Query(None),
    date_from: Optional[datetime] = Query(None),
    date_to: Optional[datetime] = Query(None),
    search: Optional[str] = Query(None),

    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user),
):
    validate_page_params(page, page_size)
    offset = (page - 1) * page_size

    query = (
        db.query(Order)
        .join(User, Order.user_id == User.id)
        .outerjoin(OrderItem, OrderItem.order_id == Order.id)
        .outerjoin(Product, Product.id == OrderItem.product_id)
        .options(
            joinedload(Order.user),
            joinedload(Order.items).joinedload(OrderItem.product),
        )
        .distinct()
    )

    # Filters
    if order_id:
        query = query.filter(Order.id == order_id)

    if user_id:
        query = query.filter(Order.user_id == user_id)

    if status:
        query = query.filter(Order.status == status)

    if payment_status:
        query = query.filter(Order.payment_status == payment_status)

    if date_from:
        query = query.filter(Order.created_at >= date_from)

    if date_to:
        query = query.filter(Order.created_at <= date_to)

    if search:
        term = f"%{search}%"
        query = query.filter(
            cast(Order.id, String).ilike(term) |
            User.first_name.ilike(term) |
            User.last_name.ilike(term) |
            User.email.ilike(term) |
            Product.name.ilike(term)
        )

    total = query.count()

    orders = (
        query
        .order_by(Order.created_at.desc())
        .offset(offset)
        .limit(page_size)
        .all()
    )

    response_orders = []

    for order in orders:
        subtotal = order.total_amount if hasattr(order, "total_amount") else 0
        gst_amount = round(subtotal * GST_RATE, 2)
        total_with_gst = round(subtotal + gst_amount, 2)

        response_orders.append({
            "order": order,
            "subtotal": subtotal,
            "gst_18_percent": gst_amount,
            "total_with_gst": total_with_gst
        })

    return {
        "gstin": "27AAMCV4451F1Z4",
        "gst_rate": "18%",
        "data": response_orders,
        "total": total
    }

# -----------------------------------------------------
# 💬 CUSTOMER SUPPORT MANAGEMENT
# -----------------------------------------------------
@router.get("/admin_support")
def list_support_messages(
    page: int = Query(DEFAULT_PAGE, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)
):
    """List all support messages (paginated)."""
    validate_page_params(page, page_size)
    offset = (page - 1) * page_size
    messages = db.query(SupportMessage).order_by(SupportMessage.created_at.desc()).offset(offset).limit(page_size).all()
    return [
        {
            "id": msg.id,
            "user_id": msg.user_id,
            "subject": msg.subject,
            "message": msg.message,
            "created_at": msg.created_at.isoformat() if getattr(msg, "created_at", None) else None
        }
        for msg in messages
    ]

@router.post("/support/reply/{message_id}")
def reply_support_message(
    message_id: int,
    payload: ReplySupportRequest = Body(...),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin_user)
):
    """Reply to a user support message (validated)."""
    # validate length was already done by Pydantic; enforce message exists
    msg = db.query(SupportMessage).filter(SupportMessage.id == message_id).first()
    if not msg:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Support message not found")

    # find user and email
    user = db.query(User).filter(User.id == msg.user_id).first()
    if not user or not getattr(user, "email", None):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Associated user not found or email missing")

    reply_text = payload.reply.strip()
    if not reply_text:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Reply cannot be empty")

    # Send email (best-effort). If sending fails, return a warning but do not crash.
    try:
        send_email(
            to_email=user.email,
            subject=f"Reply to your support ticket #{msg.id}",
            body=reply_text
        )
    except Exception as exc:
        # log exception if you have logging, here we return informative message
        return {
            "message": f"Failed to send reply email to {user.email}. Error: {str(exc)}",
            "sent": False
        }

    # Optionally: record admin reply in DB if SupportMessage model has reply fields
    if hasattr(msg, "admin_reply") and hasattr(msg, "replied_at"):
        try:
            msg.admin_reply = reply_text
            msg.replied_at = datetime.utcnow()
            db.commit()
        except Exception:
            db.rollback()

    return {"message": f"Reply sent to {user.email}", "sent": True}


